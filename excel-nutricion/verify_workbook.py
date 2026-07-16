#!/usr/bin/env python3
"""Verifica estructura mínima del libro generado."""

from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(__file__).parent / "Seguimiento_Nutricional_Deportivo.xlsx"

REQUIRED_SHEETS = [
    "Inicio",
    "Deportistas",
    "Config_ISAK",
    "Captura_ISAK",
    "Seguimiento",
    "Informe",
    "Auditoria",
    "Listas",
]

REQUIRED_TABLES = [
    "TablaDeportistas",
    "TablaConfigISAK",
    "TablaSeguimiento",
    "TablaAuditoria",
]


def main():
    assert XLSX.exists(), f"No existe {XLSX}"
    wb = load_workbook(XLSX)
    for sheet in REQUIRED_SHEETS:
        assert sheet in wb.sheetnames, f"Falta hoja {sheet}"
    for name in REQUIRED_TABLES:
        found = any(name in ws.tables for ws in wb.worksheets)
        assert found, f"Falta tabla {name}"

    dep = wb["Deportistas"]
    seg = wb["Seguimiento"]
    assert dep.max_row >= 200, "Deportistas debe tener ~200 filas"
    assert seg.max_row >= 100, "Seguimiento debe tener filas históricas"

    sample_rows = sum(
        1
        for r in range(2, 20)
        if seg.cell(r, 1).value and seg.cell(r, 3).value
    )
    assert sample_rows >= 5, "Deben existir evaluaciones de muestra"

    informe = wb["Informe"]
    assert informe.print_area, "Informe debe tener área de impresión"

    print("OK: libro verificado correctamente")


if __name__ == "__main__":
    main()
